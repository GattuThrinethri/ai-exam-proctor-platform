import io
import csv
import pytest
from httpx import AsyncClient
import docx
import openpyxl
from pptx import Presentation
import pypdf

async def get_auth_headers(client: AsyncClient, email: str, password: str) -> dict:
    """Helper to log in and return authorization headers."""
    res = await client.post("/api/auth/login", json={"email": email, "password": password})
    assert res.status_code == 200, f"Login failed: {res.text}"
    token = res.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}

def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Operating Systems Examination Questions", 0)
    doc.add_paragraph("Subject: Operating Systems")
    doc.add_paragraph("1. Explain virtual memory management and paging in detail. [10 marks]")
    doc.add_paragraph("Explanation: Paging divides memory into fixed-size frames.")
    doc.add_paragraph("2. What is a semaphore in concurrent programming?")
    doc.add_paragraph("3. True or False: Round Robin is a non-preemptive scheduling algorithm.")
    doc.add_paragraph("Answer: False")
    
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()

def create_sample_pdf() -> bytes:
    writer = pypdf.PdfWriter()
    # Add a blank page with annotation or use pypdf page
    page = writer.add_blank_page(width=612, height=792)
    # pypdf writer blank page doesn't have text unless drawn, so for text extraction test:
    # We can write an uncompressed PDF stream containing text
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()

def create_sample_pptx() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Subject: Computer Networks"
    content = slide.shapes.placeholders[1]
    content.text = "1. Which layer of the OSI model does IP reside in?\nA. Transport\nB. Network\nC. Data Link\nD. Application\nAnswer: B"
    
    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()

def create_sample_csv() -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Question", "Option A", "Option B", "Option C", "Option D", "Answer", "Marks", "Subject"])
    writer.writerow(["What is SQL used for?", "Stylesheets", "Database Queries", "Markup", "Compilation", "B", "2", "DBMS"])
    writer.writerow(["Which protocol is connection-oriented?", "UDP", "TCP", "IP", "DNS", "B", "2", "DBMS"])
    return output.getvalue().encode("utf-8")

def create_sample_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Question", "Option A", "Option B", "Option C", "Option D", "Answer", "Marks", "Subject"])
    ws.append(["What is the capital of France?", "Berlin", "Madrid", "Paris", "Rome", "C", "1", "General Knowledge"])
    
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()

@pytest.mark.asyncio
async def test_01_import_csv_questions(client: AsyncClient):
    """Test uploading a valid CSV question file and receiving structured preview."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    csv_bytes = create_sample_csv()

    files = {"file": ("test_questions.csv", csv_bytes, "text/csv")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 200, f"Import preview failed: {res.text}"
    data = res.json()
    assert data["filename"] == "test_questions.csv"
    assert data["file_type"] == "csv"
    assert data["detected_subject"] == "DBMS"
    assert data["total_extracted"] == 2
    assert len(data["questions"]) == 2

    q1 = data["questions"][0]
    assert "What is SQL used for?" in q1["question_text"]
    assert len(q1["options"]) == 4
    assert q1["correct_answer"] == "B"
    # Option B should be marked is_correct = True
    assert q1["options"][1]["is_correct"] is True
    assert q1["options"][0]["is_correct"] is False

@pytest.mark.asyncio
async def test_02_import_docx_questions(client: AsyncClient):
    """Test uploading a valid DOCX question file."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    docx_bytes = create_sample_docx()

    files = {"file": ("sample.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 200, f"DOCX import failed: {res.text}"
    data = res.json()
    assert data["total_extracted"] >= 2
    assert data["detected_subject"] == "Operating Systems"

@pytest.mark.asyncio
async def test_03_import_pptx_questions(client: AsyncClient):
    """Test uploading a valid PPTX question file."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    pptx_bytes = create_sample_pptx()

    files = {"file": ("slides.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 200, f"PPTX import failed: {res.text}"
    data = res.json()
    assert data["total_extracted"] >= 1
    assert data["questions"][0]["correct_answer"] == "B"

@pytest.mark.asyncio
async def test_04_import_xlsx_questions(client: AsyncClient):
    """Test uploading a valid XLSX question file."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    xlsx_bytes = create_sample_xlsx()

    files = {"file": ("questions.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 200, f"XLSX import failed: {res.text}"
    data = res.json()
    assert data["total_extracted"] == 1
    assert "Paris" in [opt["option_text"] for opt in data["questions"][0]["options"]]

@pytest.mark.asyncio
async def test_05_import_txt_questions(client: AsyncClient):
    """Test uploading a valid TXT question file."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    txt_content = """
Subject: Software Engineering

1. What does the acronym DRY stand for in software engineering?
A. Don't Repeat Yourself
B. Do Repeat Yourself
C. Data Retrieval Yield
D. Distributed Routing Yield
Answer: A

2. Explain the purpose of Continuous Integration (CI). [5 marks]
"""
    files = {"file": ("test.txt", txt_content.encode("utf-8"), "text/plain")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 200, f"TXT import failed: {res.text}"
    data = res.json()
    assert data["total_extracted"] == 2
    assert data["detected_subject"] == "Software Engineering"

@pytest.mark.asyncio
async def test_06_reject_unsupported_file_extension(client: AsyncClient):
    """Test uploading an unsupported file extension returns 400."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    files = {"file": ("test.exe", b"binary content", "application/x-msdownload")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 400
    assert "Unsupported file type" in res.json()["detail"]

@pytest.mark.asyncio
async def test_07_reject_oversized_file(client: AsyncClient):
    """Test uploading a file that exceeds the 15MB size limit."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    # 16 MB dummy payload
    oversized = b"A" * (16 * 1024 * 1024)
    files = {"file": ("oversized.txt", oversized, "text/plain")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 400
    assert "exceeds maximum allowed limit" in res.json()["detail"]

@pytest.mark.asyncio
async def test_08_empty_or_no_questions_document(client: AsyncClient):
    """Test uploading an empty or malformed document returns 400."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    files = {"file": ("empty.txt", b"", "text/plain")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 400

    # Non-empty document with zero questions
    files2 = {"file": ("no_questions.txt", b"Just some random greeting with no questions at all.", "text/plain")}
    res2 = await client.post("/api/questions/import", files=files2, headers=headers)
    assert res2.status_code == 400
    assert "No questions could be extracted" in res2.json()["detail"]

@pytest.mark.asyncio
async def test_09_questions_missing_answers_not_hallucinated(client: AsyncClient):
    """Verify that if questions have no answer in the document, answers are NOT hallucinated."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    txt_content = """
1. What is the time complexity of bubble sort in the worst case?
A. O(1)
B. O(n)
C. O(n log n)
D. O(n^2)
"""
    files = {"file": ("unanswered.txt", txt_content.encode("utf-8"), "text/plain")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 200
    data = res.json()
    q = data["questions"][0]
    assert q["correct_answer"] is None
    # All options must have is_correct = False
    assert all(not opt["is_correct"] for opt in q["options"])

@pytest.mark.asyncio
async def test_10_duplicate_detection_against_db(client: AsyncClient):
    """Verify duplicate detection identifies existing questions in DB."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    
    # First create a unique question in DB
    unique_text = "What is the primary function of an operating system kernel component?"
    create_payload = {
        "subject": "Operating Systems",
        "question_text": unique_text,
        "question_type": "short_answer",
        "difficulty": "medium",
        "marks": 3.0,
        "negative_marks": 0.0,
        "model_answer": "Hardware abstraction and resource management."
    }
    c_res = await client.post("/api/questions", json=create_payload, headers=headers)
    assert c_res.status_code == 201

    # Now upload a document containing the same question text
    dup_txt = f"""
1. {unique_text}
2. Explain the difference between process and thread in modern OS architectures.
"""
    files = {"file": ("dup_test.txt", dup_txt.encode("utf-8"), "text/plain")}
    res = await client.post("/api/questions/import", files=files, headers=headers)
    assert res.status_code == 200
    data = res.json()
    assert data["duplicate_count"] >= 1
    # Check that q1 is flagged duplicate
    q1 = data["questions"][0]
    assert q1["is_duplicate"] is True
    assert "Identical to existing Question" in q1["duplicate_reason"]

@pytest.mark.asyncio
async def test_11_confirm_import_inserts_into_question_bank(client: AsyncClient):
    """Test confirmation endpoint inserts selected questions into QuestionBank and QuestionOption."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")

    confirm_payload = {
        "questions": [
            {
                "question_text": "Which data structure follows LIFO principle?",
                "question_type": "MCQ",
                "options": [
                    {"option_text": "Queue", "is_correct": False},
                    {"option_text": "Stack", "is_correct": True},
                    {"option_text": "Tree", "is_correct": False},
                    {"option_text": "Graph", "is_correct": False}
                ],
                "subject": "Data Structures",
                "difficulty": "easy",
                "marks": 2.0,
                "negative_marks": 0.5
            },
            {
                "question_text": "Describe the difference between BFS and DFS traversal algorithms.",
                "question_type": "long_answer",
                "options": [],
                "subject": "Data Structures",
                "difficulty": "medium",
                "marks": 5.0,
                "negative_marks": 0.0,
                "model_answer": "BFS uses a queue and explores level-by-level; DFS uses a stack or recursion."
            }
        ],
        "default_subject": "Data Structures"
    }

    res = await client.post("/api/questions/import/confirm", json=confirm_payload, headers=headers)
    assert res.status_code == 201, f"Confirm failed: {res.text}"
    data = res.json()
    assert data["imported_count"] == 2
    assert len(data["question_ids"]) == 2

    # Verify questions appear in Question Bank list
    list_res = await client.get("/api/questions?subject=Data%20Structures", headers=headers)
    assert list_res.status_code == 200
    items = list_res.json()["items"]
    q_texts = [item["question_text"] for item in items]
    assert "Which data structure follows LIFO principle?" in q_texts

@pytest.mark.asyncio
async def test_12_role_authorization_guards(client: AsyncClient):
    """Students should receive 403 Forbidden; unauthenticated receives 401."""
    # Unauthenticated
    files = {"file": ("test.txt", b"1. Question?\nA. 1\nB. 2\nAnswer: A", "text/plain")}
    res_unauth = await client.post("/api/questions/import", files=files)
    assert res_unauth.status_code in (401, 403)

    # Student login
    student_headers = await get_auth_headers(client, "student@example.com", "Student@123")
    res_student = await client.post("/api/questions/import", files=files, headers=student_headers)
    assert res_student.status_code == 403

    # Student cannot confirm import
    res_student_confirm = await client.post(
        "/api/questions/import/confirm",
        json={"questions": [{"question_text": "Forbidden?", "question_type": "short_answer"}]},
        headers=student_headers
    )
    assert res_student_confirm.status_code == 403

@pytest.mark.asyncio
async def test_13_manual_question_creation_preserved(client: AsyncClient):
    """Verify standard manual question creation continues working normally."""
    headers = await get_auth_headers(client, "examiner@example.com", "Examiner@123")
    manual_payload = {
        "subject": "Cloud Computing",
        "question_text": "What is the primary characteristic of Infrastructure as a Service (IaaS)?",
        "question_type": "MCQ",
        "difficulty": "medium",
        "marks": 2.0,
        "negative_marks": 0.0,
        "options": [
            {"option_text": "Provides virtualized computing resources over the internet", "is_correct": True},
            {"option_text": "Provides finished software application only", "is_correct": False},
            {"option_text": "Provides physical bare-metal hardware only", "is_correct": False},
            {"option_text": "Requires on-premise hardware setup", "is_correct": False}
        ]
    }
    res = await client.post("/api/questions", json=manual_payload, headers=headers)
    assert res.status_code == 201
    created = res.json()
    assert created["subject"] == "Cloud Computing"
    assert len(created["options"]) == 4
